from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).resolve().parents[1]
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(24)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
doc.core_properties.title = 'Why we chose Problems 9 and 11'
doc.core_properties.author = ''

def content(p, text):
    text = re.sub(r'\[([^]]+)\]\([^)]+\)', '', text).strip()
    for i, chunk in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        run = p.add_run(chunk)
        run.bold = bool(i % 2)

for line in (root / 'DECISION_REPORT.md').read_text(encoding='utf-8').splitlines():
    if not line.strip():
        continue
    if line.startswith('# '):
        content(doc.add_paragraph(style='Title'), line[2:])
        doc.add_paragraph('Hackathon project decision report | 17 October 2026')
    elif line.startswith('## '):
        title = line[3:].replace(' — ', ' ')
        p = doc.add_paragraph(style='Heading 1')
        if title.startswith('Problem 9') or title.startswith('How likely'):
            p.paragraph_format.page_break_before = True
        content(p, title)
    elif line.startswith('- '):
        content(doc.add_paragraph(style='List Bullet'), line[2:])
    elif line.startswith('|'):
        cells = [c.strip() for c in line.strip('|').split('|')]
        if cells[0] == 'Outcome' or re.match(r'^-+$', cells[0]):
            continue
        content(doc.add_paragraph(style='List Bullet'), '**' + cells[0] + ':** ' + cells[1])
    else:
        content(doc.add_paragraph(), line)

doc.add_heading('Sources', level=1)
doc.add_paragraph('Official InnoSmart event information and portal documents, reviewed on 17 October 2026. The practical recommendations and feasibility assessment are our judgment, not organizer guarantees.')
doc.add_paragraph('Event page: https://innosmart.in/ics/2026/MegaHackathon-Oct26/')
doc.add_paragraph('Problem statements: H4 The Medication Safety Second Pair of Eyes; Smart Data Masking System for PII Protection. Available through the competition portal: https://app.innosmart.in/rp/competitions/13')
out = root / 'Hackathon Decision Report.docx'
for el in [doc.styles.element, doc.element]:
    for border in el.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(out)
print(out)
