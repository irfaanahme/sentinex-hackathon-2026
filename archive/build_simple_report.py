from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).resolve().parents[1]
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.75)
sec.left_margin = sec.right_margin = Inches(.85)
for name in ['Normal', 'Title', 'Heading 1', 'List Bullet']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(12)
doc.styles['Normal'].paragraph_format.space_after = Pt(8)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(23)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
doc.add_paragraph('Our hackathon choices', 'Title')
doc.add_paragraph('We chose Problem 9 and Problem 11 because both solve a clear problem, can be demonstrated simply, and give us practical ways to check whether our solution works.')

def para(label, text):
    p=doc.add_paragraph()
    p.add_run(label+' ').bold=True
    p.add_run(text)

doc.add_heading('Problem 9  Medication Safety', level=1)
para('The idea:', 'A tool that checks a patient’s medicine list and warns the doctor about possible problems, such as two medicines that should not be taken together.')
para('How we would solve it:', 'Use the medicine reference tables provided for the competition, identify the medicines, and show the most serious warnings first with a clear explanation. The doctor makes the final decision.')
para('Why we chose it:', 'The benefit is easy to understand, and we can test the warnings against the supplied references. The challenge is recognizing misspelled medicine names and avoiding unnecessary warnings.')

doc.add_heading('Problem 11  Smart Data Masking', level=1)
para('The idea:', 'A tool that hides sensitive customer information from people who do not need to see it. For example, support staff might see only part of a card number.')
para('How we would solve it:', 'Automatically find sensitive details and show different views for different staff roles. We would demonstrate it on two databases without changing the original records.')
para('Why we chose it:', 'We can clearly show who is allowed to see what and test whether the restrictions work. The challenge is making sure hidden information cannot still be accessed another way.')

doc.add_heading('Why we preferred these over the alternatives', level=1)
para('Problem 7 — Discharge Care Plans:', 'Useful, but simplifying and translating hospital instructions must preserve every medicine, dose and date. We prefer Problem 9 because reference-based checks give us a more direct way to verify many results.')
para('Problem 2 — Returns Fraud:', 'Also useful, but deciding whether a customer complaint is genuine involves uncertain patterns and the risk of wrongly accusing someone. We prefer the more clearly defined checks in Problems 9 and 11.')

doc.add_heading('What makes this a good plan', level=1)
doc.add_paragraph('We can focus on a small working solution, test it carefully, and explain it clearly. Neither problem is easy or a guaranteed winner. Our aim is to show reliable results and a convincing demonstration rather than attempt too many features.')
for el in [doc.styles.element, doc.element]:
    for border in el.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.core_properties.title='Our hackathon choices'
doc.core_properties.author=''
out=root/'Our Hackathon Choices.docx'
doc.save(out)
print(out)
